from openpyxl import Workbook as XlsxWorkbook
from openpyxl.worksheet.table import Table as XlsxTable

from xl2tableau.engine import Workbook
from xl2tableau.formula.parser import parse
from xl2tableau.tables import detect_tables, infer_dtype, to_r1c1


def test_r1c1_signature_is_row_invariant():
    a = to_r1c1(parse("=C2*D2+$H$1"), 2, 5)
    b = to_r1c1(parse("=C9*D9+$H$1"), 9, 5)
    assert a == b == "((RC[-2]*RC[-1])+R1C8)"


def test_r1c1_distinguishes_offsets():
    assert to_r1c1(parse("=A1+B2"), 2, 2) != to_r1c1(parse("=A2+B2"), 2, 2)


def test_infer_dtype():
    assert infer_dtype([1.0, 2.0, None]) == ("integer", True)
    assert infer_dtype([1.5, 2.0]) == ("real", False)
    assert infer_dtype(["a", 1.0]) == ("string", False)
    assert infer_dtype([True, False]) == ("boolean", False)


def _wb():
    wb = Workbook()
    wb.set("S", 1, 2, "note: header starts below")
    for j, h in enumerate(["Item", "Qty", "Price", "Total"], 1):
        wb.set("S", 3, j, h)
    for r in range(4, 9):
        wb.set("S", r, 1, f"item{r}")
        wb.set("S", r, 2, r)
        wb.set("S", r, 3, 1.5)
        wb.set("S", r, 4, f"=B{r}*C{r}")
    wb.set("S", 6, 4, 99)            # hardcoded override
    wb.set("S", 7, 4, "=B7*C7+1")    # edited formula
    wb.recalc()
    return wb


def test_detects_region_and_columns():
    (t,) = detect_tables(_wb())
    assert (t.header_row, t.first_row, t.last_row, t.first_col, t.last_col) == (3, 4, 8, 1, 4)
    kinds = [(c.name, c.kind, c.dtype) for c in t.columns]
    assert kinds == [("Item", "data", "string"), ("Qty", "data", "integer"),
                     ("Price", "data", "real"), ("Total", "calc", "real")]


def test_inconsistent_formulas_are_flagged():
    (t,) = detect_tables(_wb())
    total = t.columns[3]
    assert total.signature == "(RC[-2]*RC[-1])"
    assert total.inconsistent_rows == [6, 7]
    assert total.formula == "=B4*C4"


def test_excel_table_objects_take_priority(tmp_path):
    x = XlsxWorkbook()
    ws = x.active
    ws.title = "Data"
    ws.append(["junk", "row"])
    ws.append([None])
    ws.append(["k", "v", "w"])
    ws.append(["a", 1, "=B4*2"])
    ws.append(["b", 2, "=B5*2"])
    ws.add_table(XlsxTable(displayName="Orders", ref="A3:C5"))
    p = tmp_path / "t.xlsx"
    x.save(p)
    wb = Workbook.load(str(p))
    wb.recalc()
    (t,) = detect_tables(wb)
    assert t.name == "Orders" and t.first_row == 4 and t.last_row == 5
    assert t.columns[2].kind == "calc"
