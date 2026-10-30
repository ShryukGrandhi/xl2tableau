"""Generate a realistic sales-model workbook plus independently computed expected values.

The expected values are computed directly in Python (not through the formula
engine), so they serve as ground truth for engine accuracy checks.
"""
from __future__ import annotations

import math
import random
from collections import defaultdict
from typing import Dict, List, Optional, Tuple

REGIONS = ["East", "West", "North", "South", "Central"]
PRODUCTS = [("Widget", 12.5), ("Gadget", 40.0), ("Gizmo", 7.25), ("Doohickey", 99.0), ("Sprocket", 3.1)]
TAX_RATE = 0.0825
THRESHOLD = 1000.0

HEADERS = ["Region", "Product", "Units", "Unit Price", "Discount", "Revenue", "Tax", "Tier",
           "Share", "Region Revenue", "Region Share", "Label", "Margin", "Code",
           "Region Product Units", "List Price", "Running Total"]


def _round_half_away(x: float, d: int) -> float:
    from .formula.values import excel_round

    return excel_round(x, d)


def generate_rows(n: int, seed: int = 7) -> List[Tuple]:
    rnd = random.Random(seed)
    rows = []
    for _ in range(n):
        region = rnd.choice(REGIONS)
        product, list_price = rnd.choice(PRODUCTS)
        units = float(rnd.randint(0, 60))
        price = round(list_price * rnd.uniform(0.85, 1.15), 2)
        discount = rnd.choice([None, 0.0, 0.05, 0.1, 0.15])
        rows.append((region, product, units, price, discount))
    return rows


def sales_formulas(r: int, last: int) -> List[str]:
    """Formulas for columns F..Q in row r."""
    rng = lambda c: f"${c}$2:${c}${last}"  # noqa: E731
    return [
        f"=C{r}*D{r}*(1-E{r})",
        f"=F{r}*TaxRate",
        f'=IF(F{r}>=Assumptions!$B$2,"High",IF(F{r}>=Assumptions!$B$2/5,"Mid","Low"))',
        f"=F{r}/SUM({rng('F')})",
        f"=SUMIF({rng('A')},A{r},{rng('F')})",
        f"=F{r}/J{r}",
        f'=A{r}&"-"&B{r}&" x"&C{r}',
        f"=IFERROR(ROUND((F{r}-G{r})/F{r},4),0)",
        f"=UPPER(LEFT(B{r},3))&MOD(C{r},7)",
        f"=SUMIFS({rng('C')},{rng('A')},A{r},{rng('B')},B{r})",
        f"=VLOOKUP(B{r},Products!$A$2:$B${len(PRODUCTS) + 1},2,FALSE)",
        f"=F{r}" if r == 2 else f"=Q{r - 1}+F{r}",
    ]


def expected_values(rows: List[Tuple]) -> List[Dict[str, object]]:
    """Ground truth for every formula column, computed without the engine."""
    revenue = [u * p * (1 - (d or 0.0)) for (_, _, u, p, d) in rows]
    total = math.fsum(revenue)
    by_region = defaultdict(list)
    by_rp = defaultdict(float)
    for (reg, prod, u, _, _), rev in zip(rows, revenue):
        by_region[reg.lower()].append(rev)
        by_rp[(reg.lower(), prod.lower())] += u
    region_total = {k: math.fsum(v) for k, v in by_region.items()}
    lp = {p.lower(): v for p, v in PRODUCTS}
    out, running = [], 0.0
    for (reg, prod, u, p, d), rev in zip(rows, revenue):
        tax = rev * TAX_RATE
        tier = "High" if rev >= THRESHOLD else ("Mid" if rev >= THRESHOLD / 5 else "Low")
        rr = region_total[reg.lower()]
        margin = 0.0 if rev == 0 else _round_half_away((rev - tax) / rev, 4)
        running += rev
        units_txt = str(int(u))
        out.append({
            "Revenue": rev, "Tax": tax, "Tier": tier, "Share": rev / total,
            "Region Revenue": rr, "Region Share": None if rr == 0 else rev / rr,
            "Label": f"{reg}-{prod} x{units_txt}", "Margin": margin,
            "Code": prod[:3].upper() + str(int(u % 7)),
            "Region Product Units": by_rp[(reg.lower(), prod.lower())],
            "List Price": lp[prod.lower()], "Running Total": running,
        })
    return out


def build_sales_workbook(n: int, seed: int = 7, override_row: Optional[int] = None, as_engine: bool = False):
    """Build the sample model either as an openpyxl workbook or directly as an engine Workbook.

    ``override_row`` hardcodes one Tax cell, the kind of silent manual edit the
    pattern check is meant to catch.
    """
    rows = generate_rows(n, seed)
    last = n + 1
    cells: Dict[str, Dict[Tuple[int, int], object]] = {"Sales": {}, "Products": {}, "Assumptions": {}}
    for j, h in enumerate(HEADERS, 1):
        cells["Sales"][(1, j)] = h
    for i, rec in enumerate(rows):
        r = i + 2
        for j, v in enumerate(rec, 1):
            if v is not None:
                cells["Sales"][(r, j)] = v
        for j, f in enumerate(sales_formulas(r, last), 6):
            cells["Sales"][(r, j)] = f
    if override_row is not None:
        cells["Sales"][(override_row, 7)] = 12.5
    cells["Products"][(1, 1)] = "Product"
    cells["Products"][(1, 2)] = "List Price"
    for i, (p, v) in enumerate(PRODUCTS, 2):
        cells["Products"][(i, 1)] = p
        cells["Products"][(i, 2)] = v
    cells["Assumptions"].update({(1, 1): "Tax Rate", (1, 2): TAX_RATE,
                                 (2, 1): "High Threshold", (2, 2): THRESHOLD})

    if as_engine:
        from .engine import Workbook
        from .formula.parser import parse

        wb = Workbook()
        for sheet, data in cells.items():
            for (r, c), v in data.items():
                wb.set(sheet, r, c, v)
        wb.defined_names["TAXRATE"] = parse("Assumptions!$B$1")
        return wb, rows

    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    x = openpyxl.Workbook()
    x.remove(x.active)
    for sheet, data in cells.items():
        ws = x.create_sheet(sheet)
        for (r, c), v in data.items():
            ws.cell(r, c, v)
    dn = DefinedName("TaxRate", attr_text="Assumptions!$B$1")
    if hasattr(x.defined_names, "add"):
        x.defined_names.add(dn)
    else:
        x.defined_names["TaxRate"] = dn
    return x, rows
