"""Excel audit report for a conversion.

The Summary sheet uses live Excel formulas over the Fields sheet (COUNTIF,
SUM, IFERROR), so the totals stay correct if a reviewer edits statuses.
"""
from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl import Workbook as XlsxWorkbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as XlsxTable, TableStyleInfo

from .tableau.calc import render

if TYPE_CHECKING:
    from .convert import ConversionResult

_BOLD = Font(bold=True)
_FILL = {"translated": "C6EFCE", "manual": "FFC7CE", "mismatch": "FFEB9C"}


def _excel_safe(v):
    if v is None:
        return None
    if isinstance(v, (int, float, bool)):
        return v
    s = str(v)
    return "'" + s if s.startswith("=") else s


def write_audit(result: "ConversionResult", path: str) -> None:
    x = XlsxWorkbook()
    summary = x.active
    summary.title = "Summary"
    fields = x.create_sheet("Fields")
    params = x.create_sheet("Parameters")
    incons = x.create_sheet("Inconsistencies")
    mism = x.create_sheet("Mismatches")

    checks = {(c.table, c.field): c for c in result.checks}
    fields.append(["Table", "Column", "Status", "Excel formula", "R1C1 pattern",
                   "Tableau calculation", "Rows checked", "Rows matching", "Reason / notes"])
    for table in result.tables:
        for tr in result.translations[table.name]:
            chk = checks.get((table.name, tr.column.name))
            if tr.ok:
                status = "translated" if chk is None or chk.ok else "mismatch"
                calc = render(tr.tnode)
                note = "; ".join(tr.notes + ([chk.error] if chk and chk.error else []))
            else:
                status, calc, note = "manual", None, tr.reason
            fields.append([table.name, tr.column.name, status, _excel_safe(tr.column.formula),
                           tr.column.signature, calc, chk.rows if chk else None,
                           chk.matched if chk else None, note])
            fields.cell(fields.max_row, 3).fill = PatternFill("solid", fgColor=_FILL[status])
            if chk:
                for r, xv, tv in chk.mismatches:
                    mism.append([table.name, tr.column.name, r, _excel_safe(str(xv)), _excel_safe(repr(tv))])
            for r in tr.column.inconsistent_rows[:200]:
                f = table and result.workbook.sheets[table.sheet].formulas.get((r, tr.column.col))
                incons.append([table.name, tr.column.name, r, _excel_safe(f) if f else "(constant)"])
    n = fields.max_row
    if n > 1:
        tab = XlsxTable(displayName="FieldAudit", ref=f"A1:I{n}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        fields.add_table(tab)
    for col, width in zip("ABCDEFGHI", (14, 18, 12, 40, 34, 60, 13, 14, 60)):
        fields.column_dimensions[col].width = width

    params.append(["Parameter", "Source cell", "Value"])
    for p in result.params:
        params.append([p.name, f"{p.sheet}!{get_column_letter(p.col)}{p.row}", p.value])
    incons.insert_rows(1)
    for i, h in enumerate(["Table", "Column", "Row", "Formula that differs from the column pattern"], 1):
        incons.cell(1, i, h)
    mism.insert_rows(1)
    for i, h in enumerate(["Table", "Column", "Excel row", "Excel value", "Tableau value"], 1):
        mism.cell(1, i, h)

    rows = [
        ("Calculated columns", f"=COUNTA(Fields!C2:C{max(n, 2)})"),
        ("Translated and verified", f'=COUNTIF(Fields!C2:C{max(n, 2)},"translated")'),
        ("Translated, rows disagree", f'=COUNTIF(Fields!C2:C{max(n, 2)},"mismatch")'),
        ("Needs manual work", f'=COUNTIF(Fields!C2:C{max(n, 2)},"manual")'),
        ("Coverage", "=IFERROR(B3/B2,0)"),
        ("Rows checked", f"=SUM(Fields!G2:G{max(n, 2)})"),
        ("Rows matching", f"=SUM(Fields!H2:H{max(n, 2)})"),
        ("Row agreement", "=IFERROR(B8/B7,0)"),
        ("Parameters", f"=COUNTA(Parameters!A2:A{max(params.max_row, 2)})"),
        ("Inconsistent formula cells", f"=COUNTA(Inconsistencies!A2:A{max(incons.max_row, 2)})"),
    ]
    summary.append(["Metric", "Value"])
    for label, formula in rows:
        summary.append([label, formula])
    summary["B6"].number_format = "0.0%"
    summary["B9"].number_format = "0.0%"
    summary.column_dimensions["A"].width = 30
    for ws in (summary, fields, params, incons, mism):
        for c in ws[1]:
            c.font = _BOLD
        ws.freeze_panes = "A2"
    x.save(path)
