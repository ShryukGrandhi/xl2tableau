"""Workbook model and recalculation engine.

Recalculation builds a dependency graph over formula cells and evaluates them
in topological order (Kahn's algorithm), so long chains such as running totals
never hit Python's recursion limit. A range reference is a single graph node
that depends on the formula cells inside it; formula cells depend on the range
node. This keeps the edge count O(cells + distinct ranges) instead of
O(cells x range size) when every row references a whole column.
"""
from __future__ import annotations

import bisect
import time
from collections import defaultdict, deque
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from .formula.ast import (
    MAX_ROW, Binary, Bool, Err, Func, Name, Num, Percent, Ref, Str, Unary, walk,
)
from .formula.functions import (
    CTX_FUNCTIONS, FUNCTIONS, MEMO_AGGREGATES, RANGE_ARG_FUNCTIONS, SUPPORTED, ArgError,
)
from .formula.functions import _pow
from .formula.parser import FormulaParseError, parse
from .formula.values import (
    CIRC, DIV0, NAME, VALUE, RangeVal, XlError, compare, to_bool, to_number, to_text,
)

CellKey = Tuple[str, int, int]  # (sheet, row, col)


@dataclass
class Sheet:
    name: str
    constants: Dict[Tuple[int, int], object] = field(default_factory=dict)
    formulas: Dict[Tuple[int, int], str] = field(default_factory=dict)

    @property
    def max_row(self) -> int:
        rows = [r for r, _ in self.constants] + [r for r, _ in self.formulas]
        return max(rows, default=0)

    @property
    def max_col(self) -> int:
        cols = [c for _, c in self.constants] + [c for _, c in self.formulas]
        return max(cols, default=0)


@dataclass
class RecalcStats:
    formulas: int
    parse_seconds: float
    graph_seconds: float
    eval_seconds: float
    cycle_cells: int
    parse_errors: int


class EvalContext:
    def __init__(self, use_index: bool):
        self.use_index = use_index
        self.cache: Dict = {}


class Workbook:
    def __init__(self):
        self.sheets: Dict[str, Sheet] = {}
        self.defined_names: Dict[str, Ref] = {}
        # Excel "Format as Table" objects: (sheet, display name, A1 range)
        self.excel_tables: List[Tuple[str, str, str]] = []
        self.values: Dict[CellKey, object] = {}
        self.asts: Dict[CellKey, object] = {}
        self.parse_errors: Dict[CellKey, str] = {}
        self.cycle_cells: List[CellKey] = []
        self._extent: Dict[str, int] = {}

    # ------------------------------------------------------------ building
    def sheet(self, name: str) -> Sheet:
        if name not in self.sheets:
            self.sheets[name] = Sheet(name)
        return self.sheets[name]

    def set(self, sheet: str, row: int, col: int, value) -> None:
        """Set a constant, or a formula if ``value`` is a string starting with '='."""
        sh = self.sheet(sheet)
        sh.constants.pop((row, col), None)
        sh.formulas.pop((row, col), None)
        if isinstance(value, str) and value.startswith("=") and len(value) > 1:
            sh.formulas[(row, col)] = value
        elif value is not None and value != "":
            if isinstance(value, int) and not isinstance(value, bool):
                value = float(value)
            sh.constants[(row, col)] = value

    @classmethod
    def from_openpyxl(cls, wb) -> "Workbook":
        """Build from an openpyxl workbook loaded with data_only=False."""
        from openpyxl.worksheet.formula import ArrayFormula

        out = cls()
        for ws in wb.worksheets:
            out.sheet(ws.title)
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, ArrayFormula):
                        v = v.text
                    if v is None:
                        continue
                    if hasattr(v, "isoformat"):
                        v = v.isoformat()  # dates are kept as ISO text; see README limitations
                    out.set(ws.title, cell.row, cell.column, v)
            for t in ws.tables.values():
                out.excel_tables.append((ws.title, t.displayName, t.ref))
        for name, dn in _iter_defined_names(wb):
            try:
                node = parse(dn)
            except FormulaParseError:
                continue
            if isinstance(node, Ref) and node.sheet is not None:
                out.defined_names[name.upper()] = node
        return out

    @classmethod
    def load(cls, path: str) -> "Workbook":
        import openpyxl

        return cls.from_openpyxl(openpyxl.load_workbook(path, data_only=False))

    # ------------------------------------------------------------- reading
    def get(self, sheet: str, row: int, col: int):
        key = (sheet, row, col)
        if key in self.values:
            return self.values[key]
        sh = self.sheets.get(sheet)
        if sh is None:
            return None
        return sh.constants.get((row, col))

    def resolve_name(self, name: str) -> Optional[Ref]:
        return self.defined_names.get(name.upper())

    # --------------------------------------------------------------- recalc
    def recalc(self, use_index: bool = True) -> RecalcStats:
        t0 = time.perf_counter()
        self.values.clear()
        self.asts.clear()
        self.parse_errors.clear()
        self._extent = {name: sh.max_row for name, sh in self.sheets.items()}
        for name, sh in self.sheets.items():
            for (r, c), text in sh.formulas.items():
                try:
                    self.asts[(name, r, c)] = parse(text)
                except FormulaParseError as e:
                    self.parse_errors[(name, r, c)] = str(e)
                    self.values[(name, r, c)] = NAME
        t1 = time.perf_counter()
        order, blocked = self._topo_order()
        t2 = time.perf_counter()
        ctx = EvalContext(use_index)
        for key in order:
            sheet, r, c = key
            self.values[key] = _scalar_result(self._eval(self.asts[key], sheet, r, c, ctx), r, c)
        self.cycle_cells = sorted(blocked)
        for key in blocked:
            self.values[key] = CIRC
        t3 = time.perf_counter()
        return RecalcStats(len(self.asts), t1 - t0, t2 - t1, t3 - t2, len(blocked), len(self.parse_errors))

    def _clip(self, key: Tuple) -> Tuple:
        sheet, c1, r1, c2, r2 = key
        if r2 == MAX_ROW:
            r2 = max(self._extent.get(sheet, 0), r1)
        return (sheet, c1, r1, c2, r2)

    def _ref_target(self, node, sheet: str) -> Optional[Ref]:
        if isinstance(node, Ref):
            return node
        if isinstance(node, Name):
            return self.resolve_name(node.name)
        return None

    def _topo_order(self):
        # Sorted formula rows per (sheet, col) so we can find formulas inside a range quickly.
        rows_by_col: Dict[Tuple[str, int], List[int]] = defaultdict(list)
        for (s, r, c) in self.asts:
            rows_by_col[(s, c)].append(r)
        for v in rows_by_col.values():
            v.sort()

        dependents: Dict[object, List[object]] = defaultdict(list)
        indegree: Dict[object, int] = {k: 0 for k in self.asts}
        range_nodes: Dict[Tuple, bool] = {}

        def add_edge(src, dst):
            dependents[src].append(dst)
            indegree[dst] = indegree.get(dst, 0) + 1

        for key, ast in self.asts.items():
            sheet = key[0]
            seen = set()
            for node in walk(ast):
                ref = self._ref_target(node, sheet)
                if ref is None:
                    continue
                if ref.is_cell:
                    tgt = (ref.sheet or sheet, ref.row1, ref.col1)
                    if tgt in self.asts and tgt not in seen:
                        seen.add(tgt)
                        add_edge(tgt, key)
                    continue
                rk = ("R",) + self._clip(ref.key(sheet))
                if rk not in range_nodes:
                    _, s, c1, r1, c2, r2 = rk
                    members = []
                    for c in range(c1, c2 + 1):
                        rows = rows_by_col.get((s, c))
                        if rows:
                            lo = bisect.bisect_left(rows, r1)
                            hi = bisect.bisect_right(rows, r2)
                            members.extend((s, r, c) for r in rows[lo:hi])
                    range_nodes[rk] = bool(members)
                    if members:
                        indegree.setdefault(rk, 0)
                        for m in members:
                            add_edge(m, rk)
                if range_nodes[rk] and rk not in seen:
                    seen.add(rk)
                    add_edge(rk, key)

        queue = deque(k for k, d in indegree.items() if d == 0)
        order: List[CellKey] = []
        while queue:
            n = queue.popleft()
            if n[0] != "R":
                order.append(n)
            for d in dependents.get(n, ()):
                indegree[d] -= 1
                if indegree[d] == 0:
                    queue.append(d)
        done = set(order)
        blocked = [k for k in self.asts if k not in done]
        return order, blocked

    # ------------------------------------------------------------- evaluate
    def range_value(self, ref: Ref, sheet: str, ctx: EvalContext) -> RangeVal:
        key = self._clip(ref.key(sheet))
        cached = ctx.cache.get(("range", key))
        if cached is not None:
            return cached
        s, c1, r1, c2, r2 = key
        rows = [[self.get(s, r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]
        rv = RangeVal(key, rows, r1, c1)
        ctx.cache[("range", key)] = rv
        return rv

    def evaluate_formula(self, text: str, sheet: str, row: int = 0, col: int = 0):
        """Evaluate a one-off formula against current values (call recalc first)."""
        return _scalar_result(self._eval(parse(text), sheet, row, col, EvalContext(True)), row, col)

    def _eval(self, node, sheet: str, row: int, col: int, ctx: EvalContext):
        ev = self._eval
        if isinstance(node, Num):
            return node.value
        if isinstance(node, Str):
            return node.value
        if isinstance(node, Bool):
            return node.value
        if isinstance(node, Err):
            return XlError(node.code)
        if isinstance(node, (Ref, Name)):
            ref = self._ref_target(node, sheet)
            if ref is None:
                return NAME
            if ref.is_cell:
                return self.get(ref.sheet or sheet, ref.row1, ref.col1)
            return self.range_value(ref, sheet, ctx)
        if isinstance(node, Unary):
            v = to_number(_scalar(ev(node.operand, sheet, row, col, ctx), row, col))
            if isinstance(v, XlError):
                return v
            return -v if node.op == "-" else v
        if isinstance(node, Percent):
            v = to_number(_scalar(ev(node.operand, sheet, row, col, ctx), row, col))
            return v if isinstance(v, XlError) else v / 100.0
        if isinstance(node, Binary):
            a = _scalar(ev(node.left, sheet, row, col, ctx), row, col)
            b = _scalar(ev(node.right, sheet, row, col, ctx), row, col)
            return _binary(node.op, a, b)
        if isinstance(node, Func):
            return self._call(node, sheet, row, col, ctx)
        raise TypeError(f"unknown node {node!r}")

    def _call(self, node: Func, sheet: str, row: int, col: int, ctx: EvalContext):
        name, args = node.name, node.args
        if name not in SUPPORTED:
            return NAME
        ev = self._eval
        if name == "IF":
            if not 1 <= len(args) <= 3:
                return VALUE
            cond = to_bool(_scalar(ev(args[0], sheet, row, col, ctx), row, col))
            if isinstance(cond, XlError):
                return cond
            if cond:
                return ev(args[1], sheet, row, col, ctx) if len(args) > 1 else True
            return ev(args[2], sheet, row, col, ctx) if len(args) > 2 else False
        if name == "IFERROR":
            if len(args) != 2:
                return VALUE
            v = _scalar(ev(args[0], sheet, row, col, ctx), row, col)
            return _scalar(ev(args[1], sheet, row, col, ctx), row, col) if isinstance(v, XlError) else v
        if name in ("AND", "OR"):
            want = name == "OR"
            seen_any = False
            for a in args:
                v = ev(a, sheet, row, col, ctx)
                vals = v.flat() if isinstance(v, RangeVal) else [v]
                for x in vals:
                    if isinstance(v, RangeVal) and (x is None or isinstance(x, str)):
                        continue
                    b = to_bool(x)
                    if isinstance(b, XlError):
                        return b
                    seen_any = True
                    if b == want:
                        return want
            return (not want) if seen_any else VALUE

        memo_key = None
        if ctx.use_index and name in MEMO_AGGREGATES and args and all(
                isinstance(a, Ref) and not a.is_cell for a in args):
            memo_key = ("agg", name, tuple(self._clip(a.key(sheet)) for a in args))
            if memo_key in ctx.cache:
                return ctx.cache[memo_key]

        vals = []
        for a in args:
            v = ev(a, sheet, row, col, ctx)
            if isinstance(v, RangeVal) and name not in RANGE_ARG_FUNCTIONS:
                v = _scalar(v, row, col)
            vals.append(v)
        try:
            if name in CTX_FUNCTIONS:
                result = CTX_FUNCTIONS[name](ctx, *vals)
            else:
                result = FUNCTIONS[name](*vals)
        except ArgError as e:
            result = e.err
        except TypeError:
            result = VALUE  # wrong number of arguments
        if memo_key is not None:
            ctx.cache[memo_key] = result
        return result


def _iter_defined_names(wb):
    dn = wb.defined_names
    items = dn.items() if hasattr(dn, "items") else ((d.name, d) for d in dn.definedName)
    for name, d in items:
        if getattr(d, "localSheetId", None) is None and d.attr_text:
            yield name, d.attr_text


def _scalar(v, row: int, col: int):
    """Implicit intersection: a range used as a scalar picks the current row/column."""
    if not isinstance(v, RangeVal):
        return v
    if v.height == 1 and v.width == 1:
        return v.rows[0][0]
    if v.width == 1 and v.first_row <= row < v.first_row + v.height:
        return v.rows[row - v.first_row][0]
    if v.height == 1 and v.first_col <= col < v.first_col + v.width:
        return v.rows[0][col - v.first_col]
    return VALUE


def _scalar_result(v, row: int, col: int):
    v = _scalar(v, row, col)
    return 0.0 if v is None else v


def _binary(op: str, a, b):
    if op in ("=", "<>", "<", ">", "<=", ">="):
        return compare(op, a, b)
    if op == "&":
        ta, tb = to_text(a), to_text(b)
        if isinstance(ta, XlError):
            return ta
        if isinstance(tb, XlError):
            return tb
        return ta + tb
    x, y = to_number(a), to_number(b)
    if isinstance(x, XlError):
        return x
    if isinstance(y, XlError):
        return y
    if op == "+":
        return x + y
    if op == "-":
        return x - y
    if op == "*":
        return x * y
    if op == "/":
        return DIV0 if y == 0 else x / y
    if op == "^":
        return _pow(x, y)
    raise ValueError(op)
